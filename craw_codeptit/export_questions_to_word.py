import re
import os
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

BASE_DIR = Path(__file__).resolve().parent
HTML_PATH = Path(os.environ.get("PTIT_INPUT_HTML", str(BASE_DIR / "output.html")))
DOCX_PATH = Path(os.environ.get("PTIT_OUTPUT_DOCX", str(BASE_DIR / "output.docx")))


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def html_to_text_with_breaks(tag) -> str:
    clone = BeautifulSoup(str(tag), "html.parser")
    for break_tag in clone.find_all("br"):
        break_tag.replace_with("\n")

    text = clone.get_text(separator=" ", strip=False).replace("\xa0", " ")
    lines = [clean_text(line) for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def is_multiline_block(text: str) -> bool:
    lines = [line for line in text.splitlines() if line.strip()]
    return len(lines) > 1


def add_text_block(document: Document, text: str) -> None:
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return

    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.text = ""

    for index, line in enumerate(lines):
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(10)
        if index == 0:
            paragraph.paragraph_format.space_before = Pt(0)


def add_rendered_text(document: Document, text: str) -> None:
    if is_multiline_block(text):
        add_text_block(document, text)
    else:
        document.add_paragraph(text)


def add_html_table(document: Document, html_table) -> None:
    rows = html_table.find_all("tr")
    if not rows:
        return

    table_columns = max(len(row.find_all(["th", "td"], recursive=False)) for row in rows)
    table = document.add_table(rows=len(rows), cols=table_columns)
    table.style = "Table Grid"

    for row_index, row in enumerate(rows):
        cells = row.find_all(["th", "td"], recursive=False)
        for column_index in range(table_columns):
            cell = table.rows[row_index].cells[column_index]
            cell.text = ""
            if column_index >= len(cells):
                continue

            html_cell = cells[column_index]
            text = html_to_text_with_breaks(html_cell)
            lines = [line for line in text.splitlines() if line.strip()]
            if not lines:
                lines = [""]

            is_header_cell = row_index == 0 or html_cell.find(["strong", "b", "th"]) is not None
            for line_index, line in enumerate(lines):
                paragraph = cell.paragraphs[0] if line_index == 0 else cell.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(line)
                run.bold = is_header_cell
                run.font.size = Pt(10)
                run.font.name = "Times New Roman"


def extract_question_blocks(container, question_number: str) -> list[str | tuple[str, object]]:
    blocks: list[str | tuple[str, object]] = []
    seen_table = False

    for child in container.children:
        if getattr(child, "name", None) == "b":
            continue
        if getattr(child, "name", None) == "p":
            text = html_to_text_with_breaks(child)
            if question_number:
                text = re.sub(rf"^Câu\s*{question_number}\.\s*", "", text)
            if text:
                blocks.append(text)
        elif getattr(child, "name", None) == "table":
            blocks.append(("table", child))
            seen_table = True

    if not blocks and not seen_table:
        text = html_to_text_with_breaks(container)
        if question_number:
            text = re.sub(rf"^Câu\s*{question_number}\.\s*", "", text)
        if text:
            blocks.append(text)

    return blocks


def extract_question_rows(soup: BeautifulSoup) -> list:
    rows = []
    for row in soup.find_all("tr"):
        if not row.find("input", attrs={"name": re.compile(r"^mc_\d+$")}):
            continue
        question_header = row.find("b")
        if question_header is None:
            continue
        match = re.search(r"Câu\s*(\d+)", question_header.get_text(" ", strip=True))
        question_number = match.group(1) if match else ""

        question_container = row.find("span")
        if question_container is None:
            question_container = row

        question_blocks = extract_question_blocks(question_container, question_number)

        answers = []
        for label in row.find_all("label"):
            option_input = label.find("input")
            label_text = clean_text(label.get_text(" ", strip=True))
            match = re.match(r"([A-E])\.\s*", label_text)
            option_letter = match.group(1) if match else ""
            option_text_parts = []
            for child in label.children:
                if getattr(child, "name", None) == "input":
                    continue
                if hasattr(child, "get_text"):
                    option_text_parts.append(html_to_text_with_breaks(child))
                else:
                    option_text_parts.append(clean_text(str(child)))
            option_text = clean_text(" ".join(option_text_parts))
            is_selected = bool(option_input and option_input.has_attr("checked"))
            answers.append({"letter": option_letter, "text": option_text, "selected": is_selected})

        rows.append({"number": question_number, "question_blocks": question_blocks, "answers": answers})
    return rows


def main() -> None:
    if not HTML_PATH.exists():
        raise SystemExit(f"Missing input file: {HTML_PATH}")

    soup = BeautifulSoup(HTML_PATH.read_text(encoding="utf-8"), "html.parser")
    questions = extract_question_rows(soup)

    document = Document()
    document.add_heading("Danh sách câu hỏi và đáp án", level=0)
    document.add_paragraph(f"Nguồn: {HTML_PATH.name}")

    for item in questions:
        heading = f"Câu {item['number']}" if item["number"] else "Câu hỏi"
        document.add_heading(heading, level=1)
        for block in item["question_blocks"]:
            if isinstance(block, tuple) and block[0] == "table":
                add_html_table(document, block[1])
            else:
                add_rendered_text(document, block)

        for answer in item["answers"]:
            label = f"{answer['letter']}. {answer['text']}"
            if answer["selected"]:
                label += " (đang chọn)"
            if is_multiline_block(answer["text"]):
                document.add_paragraph(f"{answer['letter']}.")
                add_text_block(document, answer["text"])
                if answer["selected"]:
                    document.add_paragraph("(đang chọn)")
            else:
                document.add_paragraph(label, style="List Bullet")

    document.save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    main()